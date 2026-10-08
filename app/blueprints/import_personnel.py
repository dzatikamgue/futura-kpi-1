"""Import du personnel : photo / PDF / Excel / texte → Claude → vérification RH → enregistrement."""
import io
import unicodedata

from flask import (Blueprint, abort, current_app, flash, redirect,
                   render_template, request, send_file, url_for)
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func, select

from ..extensions import db
from ..models import Affectation, Departement, Employe, ImportBrouillon, Projet
from ..permissions import rh_requis
from ..services.audit import journaliser
from ..services.comptes import synchroniser_comptes
from ..services.organisation import (affecter, cle_texte, matricule_auto,
                                     poste_canonique, postes_actifs)
from ..services.reglages import cle_api
from ..services.import_ia import (CHAMPS, ENTETES_MODELE, ErreurImport,
                                  extraire_personnel, lire_tableur_modele)

bp = Blueprint("import_personnel", __name__, url_prefix="/import")


_cle = cle_texte


@bp.get("/")
@login_required
@rh_requis
def index():
    return render_template("import/index.html", ia_disponible=bool(cle_api()))


@bp.get("/modele.xlsx")
@login_required
@rh_requis
def modele():
    wb = Workbook()
    ws = wb.active
    ws.title = "Personnel"
    entetes = list(ENTETES_MODELE.values())
    for i, h in enumerate(entetes, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="213E70")
        ws.column_dimensions[c.column_letter].width = 22
    ws.append(["FUT-101", "ESSOMBA", "Martin", "Conducteur de travaux", "m.essomba@exemple.cm",
               "+237 6 99 00 00 00", "Direction technique", "Immeuble BALI R+11", "FUT-010"])
    ws.freeze_panes = "A2"
    dv = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="30",
                        error="30 caractères maximum")
    ws.add_data_validation(dv)
    dv.add("A2:A2000")
    notice = wb.create_sheet("Mode d'emploi")
    for ligne in [
        "Une ligne par salarié et par rattachement.",
        "Département et Projet : remplissez l'un, l'autre ou les deux.",
        "N+1 : matricule (de préférence) ou nom complet du supérieur direct.",
        "Les départements / projets inconnus pourront être créés lors de la validation.",
        "Aucune donnée n'est enregistrée avant votre validation sur l'écran de vérification.",
    ]:
        notice.append([ligne])
    notice.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="modele_import_personnel_futura.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.post("/analyser")
@login_required
@rh_requis
def analyser():
    fichier = request.files.get("fichier")
    texte = request.form.get("texte") or ""
    contenu, nom, mime = None, "", ""
    if fichier and fichier.filename:
        contenu = fichier.read()
        nom, mime = fichier.filename, fichier.mimetype or ""
    if not contenu and not texte.strip():
        flash("Ajoutez un fichier ou collez une liste.", "erreur")
        return redirect(url_for("import_personnel.index"))

    personnes, remarques, methode = None, "", "claude"
    # 1) Fichier au format du modèle : lecture directe, sans IA (rapide, gratuit, fiable)
    if contenu and nom.lower().endswith((".xlsx", ".xlsm", ".csv")):
        try:
            personnes = lire_tableur_modele(nom, contenu)
        except ErreurImport:
            personnes = None
        if personnes is not None:
            methode = "modele"
    # 2) Sinon : Claude
    if personnes is None:
        try:
            personnes, remarques = extraire_personnel(nom, mime, contenu, texte)
        except ErreurImport as exc:
            flash(str(exc), "erreur")
            return redirect(url_for("import_personnel.index"))
        except Exception:  # erreur imprévue : on journalise côté serveur
            current_app.logger.exception("Import Claude")
            flash("L'analyse a échoué de façon inattendue. Réessayez ou utilisez le modèle Excel.", "erreur")
            return redirect(url_for("import_personnel.index"))
    if not personnes:
        flash("Aucune personne détectée dans ce document.", "erreur")
        return redirect(url_for("import_personnel.index"))

    # On ne garde qu'un brouillon par utilisateur
    db.session.execute(db.delete(ImportBrouillon).where(ImportBrouillon.utilisateur_id == current_user.id))
    b = ImportBrouillon(utilisateur_id=current_user.id, source_nom=nom or "Texte collé",
                        donnees={"personnes": personnes, "methode": methode}, remarques=remarques)
    db.session.add(b)
    journaliser("import_analyse", b.source_nom, f"{len(personnes)} lignes ({methode})")
    db.session.commit()
    return redirect(url_for("import_personnel.verifier", brouillon_id=b.id))


def _brouillon(brouillon_id):
    b = db.session.get(ImportBrouillon, brouillon_id) or abort(404)
    if b.utilisateur_id != current_user.id:
        abort(403)
    return b


def _annoter(personnes):
    """Prépare les listes déroulantes de chaque ligne et signale ce qui va se passer."""
    matricules = {m for m in (p.get("matricule") for p in personnes) if m}
    existants = {e.matricule for e in db.session.scalars(
        select(Employe).where(Employe.matricule.in_(matricules or {"-"})))}
    tous = db.session.scalars(select(Employe).where(Employe.actif.is_(True))
                              .order_by(Employe.nom, Employe.prenom)).all()
    emp_par_cle = {}
    for e in tous:
        for k in (e.matricule.lower(), _cle(f"{e.nom} {e.prenom}"), _cle(f"{e.prenom} {e.nom}")):
            emp_par_cle.setdefault(k, e)
    import_par_cle = {}
    for p in personnes:
        lib = f"{(p.get('nom') or '').upper()} {p.get('prenom') or ''}".strip()
        for k in (_cle(f"{p.get('nom')} {p.get('prenom')}"), _cle(f"{p.get('prenom')} {p.get('nom')}"),
                  (p.get("matricule") or "").lower()):
            if k:
                import_par_cle.setdefault(k, lib)

    def index_ctx(model):
        idx = {}
        for x in db.session.scalars(select(model).where(model.actif.is_(True)).order_by(model.nom)):
            idx[_cle(x.nom)] = x.nom
            idx.setdefault(_cle(x.code), x.nom)
        return idx
    deps, projs = index_ctx(Departement), index_ctx(Projet)
    postes = {_cle(x.libelle): x.libelle for x in postes_actifs()}

    for p in personnes:
        p["_existe"] = bool(p.get("matricule") and p["matricule"] in existants)
        for champ, idx in (("poste", postes), ("departement", deps), ("projet", projs)):
            brut = (p.get(champ) or "").strip()
            p[f"_{champ}_val"] = idx.get(_cle(brut), brut)
            p[f"_{champ}_nouveau"] = bool(brut and _cle(brut) not in idx)
        brut = (p.get("n_plus_1") or "").strip()
        k = brut.lower() if brut.lower() in emp_par_cle else _cle(brut)
        if not brut:
            p["_n1_val"], p["_n1_ok"] = "", True
        elif k in emp_par_cle:
            p["_n1_val"], p["_n1_ok"] = emp_par_cle[k].matricule, True
        elif _cle(brut) in import_par_cle or brut.lower() in import_par_cle:
            p["_n1_val"], p["_n1_ok"] = import_par_cle.get(_cle(brut)) or import_par_cle[brut.lower()], True
        else:
            p["_n1_val"], p["_n1_ok"] = brut, False
    listes = {
        "postes": sorted(postes.values(), key=str.lower),
        "departements": sorted(set(deps.values()), key=str.lower),
        "projets": sorted(set(projs.values()), key=str.lower),
        "employes": tous,
        "importes": sorted(set(import_par_cle.values()), key=str.lower),
    }
    return personnes, listes


@bp.get("/<int:brouillon_id>/verifier")
@login_required
@rh_requis
def verifier(brouillon_id):
    b = _brouillon(brouillon_id)
    personnes, listes = _annoter([dict(p) for p in b.donnees.get("personnes", [])])
    return render_template("import/verifier.html", b=b, personnes=personnes, champs=CHAMPS,
                           methode=b.donnees.get("methode", "claude"), **listes)


def _trouver_ou_creer(model, libelle, creer, cache):
    k = _cle(libelle)
    if not k:
        return None
    if k in cache:
        return cache[k]
    obj = db.session.scalar(select(model).where(or_ci(model, libelle)))
    if not obj and creer:
        base = "".join(c for c in unicodedata.normalize("NFKD", libelle.upper()) if c.isalnum())[:8] or "CTX"
        code, i = base, 2
        while db.session.scalar(select(model.id).where(model.code == code)):
            code, i = f"{base[:6]}{i}", i + 1
        obj = model(code=code, nom=libelle.strip())
        db.session.add(obj)
        db.session.flush()
        journaliser(f"{model.__tablename__[:-1]}_cree", obj.nom, "via import")
    cache[k] = obj
    return obj


def or_ci(model, libelle):
    from sqlalchemy import or_
    v = libelle.strip().lower()
    return or_(func.lower(model.nom) == v, func.lower(model.code) == v)


@bp.post("/<int:brouillon_id>/valider")
@login_required
@rh_requis
def valider(brouillon_id):
    b = _brouillon(brouillon_id)
    n = request.form.get("nb_lignes", type=int) or 0
    creer_ctx = request.form.get("creer_contextes") == "1"
    maj = request.form.get("mettre_a_jour") == "1"

    lignes = []
    for i in range(n):
        if request.form.get(f"r{i}_inclure") != "1":
            continue
        p = {k: (request.form.get(f"r{i}_{k}") or "").strip() for k in CHAMPS}
        if not p["nom"]:
            continue
        lignes.append(p)
    if not lignes:
        flash("Aucune ligne sélectionnée.", "erreur")
        return redirect(url_for("import_personnel.verifier", brouillon_id=b.id))

    stats = {"crees": 0, "maj": 0, "ignores": 0, "affectations": 0, "n1_non_trouves": 0}
    utilises, cache_dep, cache_proj = set(), {}, {}
    employes_lignes = []
    try:
        # Passe 1 : salariés
        for p in lignes:
            e = db.session.scalar(select(Employe).where(Employe.matricule == p["matricule"])) if p["matricule"] else None
            if e and not maj:
                stats["ignores"] += 1
            elif e:
                e.nom, e.prenom = p["nom"].upper(), p["prenom"]
                e.poste = poste_canonique(p["poste"], creer_ctx) or e.poste
                e.email = (p["email"].lower() or e.email)
                e.telephone = p["telephone"] or e.telephone
                stats["maj"] += 1
            else:
                e = Employe(matricule=p["matricule"] or matricule_auto(utilises), nom=p["nom"].upper(),
                            prenom=p["prenom"], poste=poste_canonique(p["poste"], creer_ctx),
                            email=p["email"].lower() or None, telephone=p["telephone"] or None)
                db.session.add(e)
                stats["crees"] += 1
            utilises.add(e.matricule)
            employes_lignes.append((e, p))
        db.session.flush()

        # Passe 2 : affectations et N+1 (le N+1 peut être dans le même fichier)
        tous = db.session.scalars(select(Employe)).all()
        index = {e.matricule.lower(): e for e in tous}
        index.update({_cle(f"{e.nom} {e.prenom}"): e for e in tous})
        index.update({_cle(f"{e.prenom} {e.nom}"): e for e in tous})
        for e, p in employes_lignes:
            n1 = None
            if p["n_plus_1"]:
                n1 = index.get(p["n_plus_1"].lower()) or index.get(_cle(p["n_plus_1"]))
            if p["n_plus_1"] and not n1:
                stats["n1_non_trouves"] += 1
            if n1 and n1.id == e.id:
                n1 = None
            for model, champ, libelle, cache in ((Departement, "departement_id", p["departement"], cache_dep),
                                                 (Projet, "projet_id", p["projet"], cache_proj)):
                if not libelle:
                    continue
                ctx = _trouver_ou_creer(model, libelle, creer_ctx, cache)
                if not ctx:
                    continue
                deja = db.session.scalar(select(Affectation.id).where(
                    Affectation.employe_id == e.id, getattr(Affectation, champ) == ctx.id))
                affecter(e, **{("departement" if model is Departement else "projet"): ctx},
                         evaluateur_id=n1.id if n1 else None)
                if not deja:
                    stats["affectations"] += 1
        comptes = synchroniser_comptes()
        journaliser("import_valide", b.source_nom,
                    ", ".join(f"{k}={v}" for k, v in stats.items()))
        db.session.delete(b)
        db.session.commit()
    except Exception:
        db.session.rollback()
        current_app.logger.exception("Validation import")
        flash("L'import a échoué : aucune donnée n'a été enregistrée. Vérifiez les doublons de matricule.", "erreur")
        return redirect(url_for("import_personnel.verifier", brouillon_id=b.id))

    msg = (f"Import terminé : {stats['crees']} salarié(s) créé(s), {stats['maj']} mis à jour, "
           f"{stats['affectations']} affectation(s).")
    if stats["ignores"]:
        msg += f" {stats['ignores']} matricule(s) existant(s) ignoré(s)."
    if comptes["crees"]:
        msg += (f" {comptes['crees']} compte(s) de responsable créé(s) : générez leurs accès dans "
                "Comptes & accès.")
    if stats["n1_non_trouves"]:
        msg += f" {stats['n1_non_trouves']} N+1 introuvable(s) : à compléter sur les fiches concernées."
    flash(msg, "succes")
    return redirect(url_for("personnel.liste"))


@bp.post("/<int:brouillon_id>/abandonner")
@login_required
@rh_requis
def abandonner(brouillon_id):
    b = _brouillon(brouillon_id)
    db.session.delete(b)
    db.session.commit()
    flash("Import abandonné. Rien n'a été enregistré.", "info")
    return redirect(url_for("import_personnel.index"))
