"""Import d'une liste de postes (Excel, CSV ou texte collé) dans l'écran Postes.

Aucune structure ajoutée : les postes vont dans le référentiel existant (table postes).
"""
from __future__ import annotations

import csv
import io
import re

from sqlalchemy import select

from ..extensions import db
from ..models import Employe, Poste
from .entites import cond_entite
from .organisation import cle_texte, marquer_poste_supprime, postes_actifs

ENTETES = {"poste", "postes", "libelle", "intitule", "intituledeposte", "intituleduposte", "fonction",
           "fonctions", "emploi", "metier", "titre", "designation"}
MAX_POSTES = 2000


class ErreurImportPostes(ValueError):
    pass


def _norm(t) -> str:
    return re.sub(r"[^a-z0-9]", "", cle_texte(str(t or "")))


def _propre(t) -> str:
    return " ".join(str(t or "").split())[:120]


def _depuis_lignes(lignes: list[list]) -> list[str]:
    """Choisit la colonne des postes (en-tête reconnu, sinon la 1re colonne remplie)."""
    lignes = [l for l in lignes if any(_propre(v) for v in l)]
    if not lignes:
        return []
    col, debut = None, 0
    for i, l in enumerate(lignes[:5]):
        for j, v in enumerate(l):
            if _norm(v) in ENTETES:
                col, debut = j, i + 1
                break
        if col is not None:
            break
    if col is None:
        largeur = max(len(l) for l in lignes)
        col = next((j for j in range(largeur) if any(j < len(l) and _propre(l[j]) for l in lignes)), 0)
    return [_propre(l[col]) for l in lignes[debut:] if col < len(l)]


def lire_fichier(nom: str, contenu: bytes) -> list[str]:
    nom = (nom or "").lower()
    if nom.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
        except Exception as exc:
            raise ErreurImportPostes("Fichier Excel illisible. Enregistrez-le au format .xlsx puis réessayez.") from exc
        ws = wb.worksheets[0]
        return _depuis_lignes([list(r) for _, r in zip(range(MAX_POSTES + 10), ws.iter_rows(values_only=True))])
    if nom.endswith((".csv", ".txt")):
        texte = contenu.decode("utf-8-sig", errors="replace")
        premiere = next((l for l in texte.splitlines() if l.strip()), "")
        try:
            sep = csv.Sniffer().sniff(premiere, delimiters=";,\t").delimiter
        except csv.Error:
            sep = ";"
        return _depuis_lignes(list(csv.reader(io.StringIO(texte), delimiter=sep)))
    raise ErreurImportPostes("Format non pris en charge : utilisez un fichier Excel (.xlsx), CSV ou texte.")


def lire_texte(texte: str) -> list[str]:
    """Un poste par ligne (les lignes vides et l'en-tête « Poste » sont ignorés)."""
    return [l for l in (_propre(x) for x in (texte or "").splitlines()) if l and _norm(l) not in ENTETES]


def importer(entite_id: int, libelles: list[str], utilisateur_id: int | None, ne_garder_que: bool = False) -> dict:
    """Ajoute les postes absents ; renvoie le bilan. Ne modifie aucune fiche salarié."""
    vus, propres = set(), []
    for lib in libelles:
        k = cle_texte(lib)
        if lib and k not in vus:
            vus.add(k)
            propres.append(lib)
    if not propres:
        raise ErreurImportPostes("Aucun poste trouvé dans ce fichier ou ce texte.")
    if len(propres) > MAX_POSTES:
        raise ErreurImportPostes(f"Trop de postes ({len(propres)}). Maximum : {MAX_POSTES}.")

    existants = {cle_texte(x) for x in postes_actifs(entite_id)}
    lignes = {cle_texte(p.libelle): p for p in db.session.scalars(
        select(Poste).where(cond_entite(Poste.entite_id, entite_id)))}
    ajoutes, deja = [], 0
    for lib in propres:
        k = cle_texte(lib)
        if k in existants:
            deja += 1
            continue
        p = lignes.get(k)
        if p:                       # poste archivé : réactivé
            p.actif = True
        else:
            p = Poste(libelle=lib, entite_id=entite_id)
            db.session.add(p)
            lignes[k] = p
        marquer_poste_supprime(entite_id, lib, False, utilisateur_id)
        ajoutes.append(lib)

    retires, gardes = [], []
    if ne_garder_que:
        portes = {cle_texte(x) for x in db.session.scalars(select(Employe.poste).where(
            Employe.poste.isnot(None), cond_entite(Employe.entite_id, entite_id)).distinct())}
        db.session.flush()
        for lib in postes_actifs(entite_id):
            k = cle_texte(lib)
            if k in vus:
                continue
            if k in portes:         # encore porté par des fiches : conservé
                gardes.append(lib)
                continue
            if k in lignes:
                db.session.delete(lignes[k])
            marquer_poste_supprime(entite_id, lib, True, utilisateur_id)
            retires.append(lib)
    return {"lus": len(propres), "ajoutes": ajoutes, "deja": deja, "retires": retires, "gardes": gardes}
