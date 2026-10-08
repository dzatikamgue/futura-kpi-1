"""Exports Excel (openpyxl) et PDF (reportlab) aux couleurs Futura."""
from __future__ import annotations

import io
import os
from datetime import datetime

from flask import current_app, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (Image, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)

BLEU = "213E70"
GRIS = "5E5B5C"


def _nom_fichier(base: str, ext: str) -> str:
    return f"{base}_{datetime.now():%Y%m%d_%H%M}.{ext}"


def export_excel(titre: str, entetes: list[str], lignes: list[list], nom: str,
                 colonnes_num: set[int] | None = None, total: list | None = None,
                 sous_titre: str = ""):
    colonnes_num = colonnes_num or set()
    wb = Workbook()
    ws = wb.active
    ws.title = titre[:31]
    ws["A1"] = f"{current_app.config['COMPANY_NAME']} — {titre}"
    ws["A1"].font = Font(bold=True, size=14, color=BLEU)
    ws["A2"] = sous_titre or f"Exporté le {datetime.now():%d/%m/%Y à %H:%M}"
    ws["A2"].font = Font(size=10, color=GRIS)

    entete_row = 4
    fill = PatternFill("solid", fgColor=BLEU)
    fin = Side(style="thin", color="D9DCE3")
    for c, h in enumerate(entetes, 1):
        cell = ws.cell(row=entete_row, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="right" if (c - 1) in colonnes_num else "left",
                                   vertical="center", wrap_text=True)
    for r, ligne in enumerate(lignes, entete_row + 1):
        for c, v in enumerate(ligne, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = Border(bottom=fin)
            if (c - 1) in colonnes_num:
                cell.number_format = "0.00"
                cell.alignment = Alignment(horizontal="right")
    if total:
        r = entete_row + len(lignes) + 1
        for c, v in enumerate(total, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EEF1F7")
            if (c - 1) in colonnes_num:
                cell.number_format = "0.00"
                cell.alignment = Alignment(horizontal="right")
    for c in range(1, len(entetes) + 1):
        largeur = max([len(str(entetes[c - 1]))] + [len(str(l[c - 1] or "")) for l in lignes[:500]])
        ws.column_dimensions[get_column_letter(c)].width = min(max(10, largeur + 2), 45)
    ws.freeze_panes = ws.cell(row=entete_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=_nom_fichier(nom, "xlsx"),
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _fmt(v, num):
    if v is None or v == "":
        return "—"
    if num and isinstance(v, (int, float)):
        return f"{v:.2f}".replace(".", ",")
    return str(v)


def export_pdf(titre: str, entetes: list[str], lignes: list[list], nom: str,
               colonnes_num: set[int] | None = None, total: list | None = None,
               sous_titre: str = "", paysage: bool = True):
    colonnes_num = colonnes_num or set()
    buf = io.BytesIO()
    taille = landscape(A4) if paysage else A4
    doc = SimpleDocTemplate(buf, pagesize=taille, leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=12 * mm, bottomMargin=14 * mm, title=titre,
                            author=current_app.config["COMPANY_NAME"])
    st_titre = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=15,
                              textColor=colors.HexColor("#" + BLEU), spaceAfter=2)
    st_sous = ParagraphStyle("s", fontName="Helvetica", fontSize=9, textColor=colors.HexColor("#" + GRIS))
    st_cell = ParagraphStyle("c", fontName="Helvetica", fontSize=8, leading=10)

    elements = []
    logo = os.path.join(current_app.static_folder, "img", "logo.png")
    if os.path.exists(logo):
        elements.append(Image(logo, width=42 * mm, height=16 * mm, hAlign="LEFT"))
        elements.append(Spacer(1, 4 * mm))
    elements.append(Paragraph(titre, st_titre))
    elements.append(Paragraph(sous_titre or f"Édité le {datetime.now():%d/%m/%Y à %H:%M}", st_sous))
    elements.append(Spacer(1, 5 * mm))

    data = [entetes]
    for l in lignes:
        data.append([_fmt(v, i in colonnes_num) if i in colonnes_num else Paragraph(_fmt(v, False), st_cell)
                     for i, v in enumerate(l)])
    if total:
        data.append([_fmt(v, i in colonnes_num) for i, v in enumerate(total)])

    largeur_dispo = taille[0] - 28 * mm
    n_num = len(colonnes_num)
    n_txt = len(entetes) - n_num
    w_num = min(16 * mm, largeur_dispo * 0.5 / max(n_num, 1)) if n_num else 0
    w_txt = (largeur_dispo - w_num * n_num) / max(n_txt, 1)
    widths = [w_num if i in colonnes_num else w_txt for i in range(len(entetes))]

    t = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + BLEU)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#D9DCE3")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1 if not total else -2), [colors.white, colors.HexColor("#F7F8FB")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for i in colonnes_num:
        style.append(("ALIGN", (i, 0), (i, -1), "RIGHT"))
    if total:
        style += [("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                  ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#EEF1F7"))]
    t.setStyle(TableStyle(style))
    elements.append(t)

    def pied(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#" + GRIS))
        canvas.drawString(14 * mm, 8 * mm, f"{current_app.config['COMPANY_NAME']} — Document confidentiel RH")
        canvas.drawRightString(taille[0] - 14 * mm, 8 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(elements, onFirstPage=pied, onLaterPages=pied)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=_nom_fichier(nom, "pdf"),
                     mimetype="application/pdf")
