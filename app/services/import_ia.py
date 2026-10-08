"""Extraction d'une liste de personnel à partir d'une photo, d'un PDF, d'un Excel/CSV
ou d'un texte collé, via l'API Claude.

Rien n'est écrit en base ici : le résultat est proposé à la RH, qui le corrige
et le valide sur l'écran de vérification.
"""
from __future__ import annotations

import base64
import csv
import io
import re

from flask import current_app

CHAMPS = ["matricule", "nom", "prenom", "poste", "email", "telephone",
          "departement", "projet", "n_plus_1"]

ENTETES_MODELE = {
    "matricule": "Matricule", "nom": "Nom", "prenom": "Prénom", "poste": "Poste",
    "email": "Email", "telephone": "Téléphone", "departement": "Département",
    "projet": "Projet", "n_plus_1": "N+1 (nom ou matricule)",
}

TYPES_IMAGES = {"image/jpeg", "image/png", "image/webp", "image/gif"}
EXT_IMAGES = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
              ".webp": "image/webp", ".gif": "image/gif"}

OUTIL = {
    "name": "enregistrer_personnel",
    "description": "Enregistre la liste des membres du personnel extraite du document.",
    "input_schema": {
        "type": "object",
        "properties": {
            "personnes": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "matricule": {"type": "string", "description": "Matricule s'il figure, sinon chaîne vide"},
                        "nom": {"type": "string", "description": "Nom de famille"},
                        "prenom": {"type": "string"},
                        "poste": {"type": "string", "description": "Fonction / poste"},
                        "email": {"type": "string"},
                        "telephone": {"type": "string"},
                        "departement": {"type": "string", "description": "Département ou service"},
                        "projet": {"type": "string", "description": "Projet / chantier d'affectation"},
                        "n_plus_1": {"type": "string", "description": "Nom ou matricule du supérieur direct"},
                        "confiance": {"type": "string", "enum": ["haute", "moyenne", "basse"],
                                      "description": "Fiabilité de la lecture de cette ligne"},
                        "remarque": {"type": "string", "description": "Doute de lecture éventuel, sinon vide"},
                    },
                    "required": ["nom", "prenom", "confiance"],
                },
            },
            "remarques": {"type": "string", "description": "Remarques générales sur le document"},
        },
        "required": ["personnes"],
    },
}

CONSIGNE = (
    "Tu aides le service RH de l'entreprise {entreprise} (BTP, Cameroun) à saisir son personnel. "
    "Extrais TOUTES les personnes listées dans le document fourni et appelle l'outil "
    "enregistrer_personnel. Règles :\n"
    "- Ne devine jamais une information absente : laisse le champ vide.\n"
    "- Sépare correctement nom de famille et prénom(s) (au Cameroun, le nom est souvent écrit en premier, en majuscules).\n"
    "- Recopie les noms exactement, avec accents ; n'invente pas d'email.\n"
    "- Mets confiance='basse' et une remarque si l'écriture est illisible ou ambiguë.\n"
    "- Ignore les lignes de titre, totaux, signatures."
)


class ErreurImport(Exception):
    pass


def detail_erreur_api(exc) -> str:
    """Message d'origine renvoyé par Anthropic (le plus précis disponible)."""
    msg = getattr(exc, "message", "") or ""
    body = getattr(exc, "body", None)
    if isinstance(body, dict):
        msg = (body.get("error") or {}).get("message") or body.get("message") or msg
    return str(msg or exc).strip()


def traduire_erreur_api(exc) -> str:
    """Explique une erreur Anthropic en français, avec la vraie cause."""
    code = getattr(exc, "status_code", "?")
    detail = detail_erreur_api(exc)
    bas = detail.lower()
    if "credit balance" in bas or "billing" in bas:
        return ("Le compte Anthropic n'a pas de crédit : ajoutez du crédit sur console.anthropic.com → "
                "Plans & Billing, puis réessayez (quelques minutes peuvent être nécessaires).")
    if "model" in bas and any(k in bas for k in ("not found", "does not exist", "not supported", "invalid")):
        return (f"Modèle refusé par Anthropic ({detail[:160]}). Corrigez la variable ANTHROPIC_MODEL "
                "sur Railway ou supprimez-la.")
    if code == 413 or "too large" in bas or "exceeds" in bas and ("size" in bas or "mb" in bas):
        return "Le fichier est trop volumineux pour Claude. Réduisez-le (moins de pages / image plus légère)."
    if any(k in bas for k in ("image", "pdf", "media_type", "base64", "document")):
        return f"Claude n'a pas pu lire ce fichier : {detail[:240]}"
    return f"Anthropic a refusé la demande (code {code}) : {detail[:300] or 'aucun détail fourni'}"


def _tableur_en_texte(nom: str, contenu: bytes) -> str:
    """Convertit un Excel/CSV en texte CSV (limité) pour l'envoyer à Claude."""
    if nom.lower().endswith(".csv"):
        texte = contenu.decode("utf-8-sig", errors="replace")
        return texte[:60000]
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    except Exception as exc:  # fichier corrompu ou .xls ancien
        raise ErreurImport("Fichier Excel illisible. Enregistrez-le au format .xlsx puis réessayez.") from exc
    sortie = io.StringIO()
    w = csv.writer(sortie, delimiter=";")
    for ws in wb.worksheets[:3]:
        sortie.write(f"# Feuille : {ws.title}\n")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 2000:
                break
            if any(v not in (None, "") for v in row):
                w.writerow(["" if v is None else str(v) for v in row])
    return sortie.getvalue()[:60000]


def lire_tableur_modele(nom: str, contenu: bytes) -> list[dict] | None:
    """Lecture directe (sans IA) d'un fichier au format du modèle Futura.

    Retourne None si les colonnes ne correspondent pas au modèle.
    """
    texte = _tableur_en_texte(nom, contenu)
    lignes_brutes = [l for l in texte.splitlines() if l.strip() and not l.startswith("# Feuille")]
    if not lignes_brutes:
        return None
    try:
        sep = csv.Sniffer().sniff(lignes_brutes[0], delimiters=";,\t").delimiter
    except csv.Error:
        sep = ";"
    lignes = [l for l in csv.reader(io.StringIO("\n".join(lignes_brutes)), delimiter=sep) if l]
    if not lignes:
        return None
    norm = lambda s: re.sub(r"[^a-z0-9+]", "", (s or "").lower()
                            .replace("é", "e").replace("è", "e").replace("ô", "o"))
    cibles = {norm(v): k for k, v in ENTETES_MODELE.items()}
    cibles.update({"n+1": "n_plus_1", "departement": "departement"})
    entete = [cibles.get(norm(c)) for c in lignes[0]]
    if "nom" not in entete:
        return None
    personnes = []
    for l in lignes[1:]:
        p = {k: "" for k in CHAMPS}
        for i, k in enumerate(entete):
            if k and i < len(l):
                p[k] = (l[i] or "").strip()
        if p["nom"]:
            p["confiance"] = "haute"
            p["remarque"] = ""
            personnes.append(p)
    return personnes


def extraire_personnel(nom_fichier: str, type_mime: str, contenu: bytes | None, texte: str = ""):
    """Appelle Claude et renvoie (personnes, remarques)."""
    from .reglages import cle_api
    cle = cle_api()
    if not cle:
        raise ErreurImport("La clé API Claude n'est pas configurée : renseignez-la dans "
                           "Administration RH → Clé API Claude, ou utilisez le modèle Excel.")
    import anthropic

    ext = "." + nom_fichier.rsplit(".", 1)[-1].lower() if nom_fichier and "." in nom_fichier else ""
    blocs = []
    if contenu:
        mime = type_mime if type_mime in TYPES_IMAGES else EXT_IMAGES.get(ext)
        if mime:
            blocs.append({"type": "image", "source": {"type": "base64", "media_type": mime,
                                                      "data": base64.b64encode(contenu).decode()}})
        elif ext == ".pdf" or type_mime == "application/pdf":
            blocs.append({"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                                         "data": base64.b64encode(contenu).decode()}})
        elif ext in (".xlsx", ".xlsm", ".csv"):
            blocs.append({"type": "text", "text": f"Contenu du fichier {nom_fichier} :\n{_tableur_en_texte(nom_fichier, contenu)}"})
        else:
            raise ErreurImport("Format non pris en charge. Formats acceptés : photo (JPG, PNG, WEBP), PDF, Excel (.xlsx) ou CSV.")
    if texte.strip():
        blocs.append({"type": "text", "text": f"Liste collée par l'utilisateur :\n{texte[:30000]}"})
    if not blocs:
        raise ErreurImport("Ajoutez un fichier ou collez une liste.")
    blocs.append({"type": "text", "text": CONSIGNE.format(entreprise=current_app.config["COMPANY_NAME"])})

    client = anthropic.Anthropic(api_key=cle, timeout=120.0, max_retries=2)
    modele = current_app.config["ANTHROPIC_MODEL"]

    def appeler(tool_choice, max_tokens):
        return client.messages.create(
            model=modele, max_tokens=max_tokens, tools=[OUTIL], tool_choice=tool_choice,
            messages=[{"role": "user", "content": blocs}])

    try:
        try:
            reponse = appeler({"type": "tool", "name": OUTIL["name"]}, 16000)
        except anthropic.BadRequestError as exc:
            # Certains modèles / réglages refusent l'outil forcé ou une longueur de sortie : on adapte.
            bas = detail_erreur_api(exc).lower()
            if "tool_choice" in bas or "thinking" in bas:
                current_app.logger.warning("tool_choice forcé refusé (%s) : repli en mode auto", bas[:200])
                reponse = appeler({"type": "auto"}, 16000)
            elif "max_tokens" in bas:
                current_app.logger.warning("max_tokens refusé (%s) : repli à 8000", bas[:200])
                reponse = appeler({"type": "tool", "name": OUTIL["name"]}, 8000)
            else:
                raise
    except anthropic.AuthenticationError as exc:
        raise ErreurImport("Clé API Claude invalide. Corrigez-la dans Administration RH → Clé API Claude.") from exc
    except anthropic.PermissionDeniedError as exc:
        raise ErreurImport("Cette clé n'a pas le droit d'utiliser ce modèle : " + detail_erreur_api(exc)[:200]) from exc
    except anthropic.RateLimitError as exc:
        raise ErreurImport("Le service Claude est momentanément saturé. Réessayez dans une minute.") from exc
    except anthropic.APIConnectionError as exc:
        raise ErreurImport("Connexion au service Claude impossible. Vérifiez la connexion du serveur.") from exc
    except anthropic.APIStatusError as exc:
        current_app.logger.warning("Erreur API Claude %s : %s", exc.status_code, detail_erreur_api(exc))
        raise ErreurImport(traduire_erreur_api(exc)) from exc

    for bloc in reponse.content:
        if getattr(bloc, "type", "") == "tool_use":
            data = bloc.input or {}
            personnes = []
            for p in data.get("personnes", []):
                propre = {k: str(p.get(k) or "").strip() for k in CHAMPS}
                if not propre["nom"]:
                    continue
                propre["confiance"] = p.get("confiance") or "moyenne"
                propre["remarque"] = str(p.get("remarque") or "").strip()
                personnes.append(propre)
            return personnes, str(data.get("remarques") or "")
    raise ErreurImport("Claude n'a renvoyé aucune liste exploitable. Essayez avec un document plus lisible.")
